# Sinh template ORDER_FORM_X (.xlsx, engine XLSX) + logo PNG mau + Sample JSON co logo base64
import sys, json, zlib, struct, base64, pathlib
from openpyxl import Workbook
from openpyxl.styles import Font, Alignment, Border, Side, PatternFill

out = pathlib.Path(sys.argv[1])

# ---------- logo PNG 200x110 (khong can PIL) ----------
W, H, SC = 200, 110, 8
GLYPH = {
 "L": ["10000","10000","10000","10000","10000","10000","11111"],
 "O": ["01110","10001","10001","10001","10001","10001","01110"],
 "G": ["01110","10001","10000","10111","10001","10001","01111"],
}
bg, fg, accent = (31, 78, 121), (255, 255, 255), (237, 125, 49)
px = [[bg] * W for _ in range(H)]
text = "LOGO"
tw = len(text) * 5 * SC + (len(text) - 1) * SC
x0, y0 = (W - tw) // 2, (H - 7 * SC) // 2 - 6
for i, ch in enumerate(text):
    gx = x0 + i * 6 * SC
    for r, line in enumerate(GLYPH[ch]):
        for c, bit in enumerate(line):
            if bit == "1":
                for yy in range(SC):
                    for xx in range(SC):
                        px[y0 + r * SC + yy][gx + c * SC + xx] = fg
for yy in range(H - 18, H - 10):              # vach cam duoi chu
    for xx in range(x0, x0 + tw):
        px[yy][xx] = accent
raw = b"".join(b"\x00" + bytes(v for p in row for v in p) for row in px)
def chunk(t, d): return struct.pack(">I", len(d)) + t + d + struct.pack(">I", zlib.crc32(t + d) & 0xffffffff)
png = (b"\x89PNG\r\n\x1a\n" + chunk(b"IHDR", struct.pack(">IIBBBBB", W, H, 8, 2, 0, 0, 0))
       + chunk(b"IDAT", zlib.compress(raw, 9)) + chunk(b"IEND", b""))
(out / "LOGO_DEMO.png").write_bytes(png)

# ---------- template .xlsx ----------
wb = Workbook(); ws = wb.active; ws.title = "DonHang"
F  = lambda **k: Font(name="Arial", size=k.pop("size", 10), **k)
thin = Side(style="thin"); dbl = Side(style="double")
box  = Border(left=thin, right=thin, top=thin, bottom=thin)
grid = Border(left=thin, right=thin, bottom=thin)
for col, w in zip("ABCDEFG", [6.5, 17, 37, 8, 11, 16, 20]):
    ws.column_dimensions[col].width = w
for r in (1, 2, 3, 4):
    ws.row_dimensions[r].height = 18

def put(ref, val, font=None, align=None, border=None, fill=None, fmt=None, merge=None):
    c = ws[ref]; c.value = val; c.font = font or F()
    if align: c.alignment = align
    if border: c.border = border
    if fill: c.fill = fill
    if fmt: c.number_format = fmt
    if merge: ws.merge_cells(f"{ref}:{merge}")
    return c

left, center, right = Alignment(horizontal="left", vertical="center"), \
    Alignment(horizontal="center", vertical="center", wrap_text=True), Alignment(horizontal="right", vertical="center")

# logo: neo o A1, rong 2 cot (A:B), cao 4 dong (1..4) — engine co dinh 4 dong
put("A1", "{{*image:logo@2x4}}")
put("C1", "{{company.name}}", F(size=12, bold=True), left, merge="G1")
put("C2", "Địa chỉ: {{company.address}}", align=left, merge="G2")
put("C3", "MST: {{company.taxcode}}", align=left, merge="G3")
put("C4", "Điện thoại: {{company.phone}}", align=left, merge="G4")

ws.row_dimensions[6].height = 26
put("A6", "ĐƠN ĐẶT HÀNG", F(size=16, bold=True), center, merge="G6")
put("A7", "Số: {{orderno}}", F(italic=True), center, merge="G7")

put("A9", "Khách hàng:", F(bold=True), merge="B9")
put("C9", "{{customer.name}}", merge="E9")
put("F9", "Ngày đặt:", F(bold=True))
put("G9", "{{orddate}}", align=left, fmt="dd/mm/yyyy")
put("A10", "Địa chỉ:", F(bold=True), merge="B10")
put("C10", "{{customer.address}}", merge="E10")
put("F10", "Ngày giao:", F(bold=True))
put("G10", "{{deldate}}", align=left, fmt="dd/mm/yyyy")

ws.row_dimensions[12].height = 28
hdr = PatternFill("solid", fgColor="DDEBF7")
for col, t in zip("ABCDEFG", ["STT", "Mã hàng", "Tên hàng", "ĐVT", "Số lượng", "Đơn giá", "Thành tiền"]):
    put(f"{col}12", t, F(bold=True), center, box, hdr)

# dong lap: {{#items}} o dau, {{/items}} o cuoi CUNG 1 dong
row = [("A", "{{#items}}{{@index}}", center, None), ("B", "{{matnr}}", None, None),
       ("C", "{{maktx}}", None, None), ("D", "{{uom}}", center, None),
       ("E", "{{qty}}", None, "#,##0"), ("F", "{{price}}", None, "#,##0"),
       ("G", "{{amount}}{{/items}}", None, "#,##0")]
for col, v, al, fmt in row:
    put(f"{col}13", v, align=al, border=grid, fmt=fmt)

put("A14", "Cộng tiền hàng ({{cnt:items}} dòng):", F(bold=True), right, merge="F14")
put("G14", "{{sum:items.amount}}", F(bold=True), fmt="#,##0")
put("A15", "Thuế GTGT {{vat_rate}}%:", F(bold=True), right, merge="F15")
put("G15", "{{vat_amount}}", F(bold=True), fmt="#,##0")
put("A16", "Tổng thanh toán:", F(bold=True), right, merge="F16")
put("G16", "{{grand_total}}", F(size=11, bold=True), border=Border(top=dbl), fmt="#,##0")

put("A18", "{{?remark}}Ghi chú:", F(bold=True))
put("B18", "{{remark}}{{/remark}}", merge="G18")

put("B20", "Người lập phiếu", F(bold=True), center, merge="C20")
put("F20", "Khách hàng", F(bold=True), center, merge="G20")

ws.page_setup.paperSize = ws.PAPERSIZE_A4
ws.page_setup.orientation = "portrait"
ws.sheet_properties.pageSetUpPr.fitToPage = True
ws.page_setup.fitToWidth, ws.page_setup.fitToHeight = 1, 0
xlsx = out / "ORDER_FORM_X_template.xlsx"
wb.save(xlsx)

# openpyxl ghi Target tuyet doi "/xl/worksheets/sheet1.xml"; Excel ghi tuong doi
# "worksheets/sheet1.xml". Doi ve dang Excel de engine ban cu cung doc duoc sheet.
import zipfile, io
buf = io.BytesIO()
with zipfile.ZipFile(xlsx) as zin, zipfile.ZipFile(buf, "w", zipfile.ZIP_DEFLATED) as zout:
    for item in zin.infolist():
        data = zin.read(item.filename)
        if item.filename == "xl/_rels/workbook.xml.rels":
            data = data.replace(b'Target="/xl/', b'Target="')
        zout.writestr(item, data)
xlsx.write_bytes(buf.getvalue())

# ---------- Sample JSON ----------
data = json.load(open(out / "ORDER_FORM_sample.json", encoding="utf-8"))
data["company"]["phone"] = "024 1234 5678"
data = {"logo": base64.b64encode(png).decode(), **data}
(out / "ORDER_FORM_X_sample.json").write_text(json.dumps(data, ensure_ascii=False, indent=2), encoding="utf-8")
print("png bytes", len(png), "b64 chars", len(data["logo"]))
